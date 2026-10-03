#!/usr/bin/env python3
"""Тесты исправлений: dxf_inspect (POLYLINE R12, сборка контуров из LINE/ARC, --tol, --layers),
unfold (--outer/--t/--k только для mitre), cost_calc (нормализация категорий/kind, база накладных),
hx_calc heat_balance (знак Φ, небаланс, неизвестные), pipe_calc bend_b31_3 (ASME B31.3 eq. 3c).

Эталоны посчитаны вручную (значения в комментариях).

Запуск: python tests/test_cad_cost_hx.py   (или python -m pytest tests -q)
"""
from __future__ import annotations

import importlib.util
import io
import json
import logging
import math
import random
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SK = ROOT / ".claude" / "skills"


def load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.path.insert(0, str(path.parent))
    spec.loader.exec_module(mod)
    return mod


pipe = load("pipe_calc", SK / "novaprom-piping/scripts/pipe_calc.py")
hx = load("hx_calc", SK / "novaprom-heat-exchangers/scripts/hx_calc.py")
cost = load("cost_calc", SK / "novaprom-cost-estimation/scripts/cost_calc.py")
UNFOLD = SK / "novaprom-cad/scripts/unfold.py"
try:
    import ezdxf
    logging.getLogger("ezdxf").setLevel(logging.ERROR)  # «$INSUNITS не экспортируется в R12» — не ошибка
    dxfi = load("dxf_inspect", SK / "novaprom-cad/scripts/dxf_inspect.py")
except (ImportError, SystemExit):  # ezdxf не установлен
    ezdxf = dxfi = None
try:
    import openpyxl
except ImportError:
    openpyxl = None


def step_value(rep, symbol: str) -> float:
    for s in rep.steps:
        if s.result.startswith(symbol + " ="):
            return float(s.result.split("=")[1].split()[0])
    raise KeyError(symbol)


# ---------------------------------------------------------------------------------------------
@unittest.skipIf(dxfi is None, "нужен ezdxf")
class DxfInspect(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def save(self, doc, name: str) -> str:
        p = str(self.tmp / name)
        doc.saveas(p)
        return p

    def rect_with_holes(self, bend_line=False) -> str:
        doc = ezdxf.new("R2010")
        m = doc.modelspace()
        # 1000×500 из 4 LINE в произвольном порядке, часть — в обратном направлении
        m.add_line((1000, 0), (1000, 500))
        m.add_line((0, 0), (1000, 0))
        m.add_line((0, 500), (1000, 500))
        m.add_line((0, 500), (0, 0))
        m.add_circle((200, 250), 50)
        m.add_circle((800, 250), 30)
        if bend_line:
            doc.layers.add("BEND")
            m.add_line((500, 0), (500, 500), dxfattribs={"layer": "BEND"})
        return self.save(doc, "rect.dxf")

    def test_rectangle_of_lines_with_holes(self):
        # наружный 1000·500 = 500000; отверстия π·50² = 7853.98 и π·30² = 2827.43; нетто 489318.59
        r, _ = dxfi.inspect(self.rect_with_holes(), 5)
        b = r["blank_estimate"]
        self.assertEqual(r["closed_contours"][0]["type"], "CHAIN(4)")
        self.assertAlmostEqual(r["closed_contours"][0]["perimeter_mm"], 3000.0, places=2)
        self.assertAlmostEqual(b["outer_area_mm2"], 500000, delta=0.5)
        self.assertAlmostEqual(b["holes_area_mm2"], 7853.98 + 2827.43, delta=0.5)
        self.assertAlmostEqual(b["net_area_mm2"], 489319, delta=1)
        self.assertEqual(b["pierces"], 3)
        self.assertTrue(b["reliable"])
        self.assertEqual(r["open_chains"]["count"], 0)

    def test_r12_closed_polyline_does_not_crash(self):
        doc = ezdxf.new("R12")
        m = doc.modelspace()
        m.add_polyline2d([(0, 0), (100, 0), (100, 50), (0, 50)], close=True)  # старая POLYLINE: is_closed
        m.add_polyline2d([(200, 0), (300, 0)])
        r, _ = dxfi.inspect(self.save(doc, "r12.dxf"), 5)
        self.assertEqual(r["closed_contours"][0]["type"], "POLYLINE")
        self.assertAlmostEqual(r["closed_contours"][0]["area_mm2"], 5000, places=2)
        self.assertAlmostEqual(r["closed_contours"][0]["perimeter_mm"], 300, places=2)
        self.assertEqual(r["open_chains"]["count"], 1)

    def test_line_arc_outline_is_chained(self):
        # прямоугольник 1000×500 со скруглениями R50 из 4 LINE + 4 ARC (одна дуга — в зеркальной ОСК, как
        # пишет SolidWorks): A = 1000·500 − (4 − π)·50² = 497853.98; периметр 2·900 + 2·400 + 2π·50 = 2914.16
        W, H, R = 1000, 500, 50
        doc = ezdxf.new("R2010")
        m = doc.modelspace()
        m.add_line((R, 0), (W - R, 0))
        m.add_line((W, H - R), (W, R))  # обратное направление
        m.add_line((W - R, H), (R, H))
        m.add_line((0, H - R), (0, R))
        m.add_arc((W - R, R), R, 270, 360)
        m.add_arc((W - R, H - R), R, 0, 90)
        m.add_arc((R, H - R), R, 90, 180)
        m.add_arc((-R, R), R, 270, 360, dxfattribs={"extrusion": (0, 0, -1)})  # центр (R, R) в МСК
        m.add_circle((500, 250), 100)
        r, _ = dxfi.inspect(self.save(doc, "round.dxf"), 5)
        outer = r["closed_contours"][0]
        self.assertEqual(outer["type"], "CHAIN(8)")
        self.assertAlmostEqual(outer["area_mm2"], W * H - (4 - math.pi) * R * R, delta=5)
        self.assertAlmostEqual(outer["perimeter_mm"], 2 * 900 + 2 * 400 + 2 * math.pi * R, delta=0.05)
        self.assertAlmostEqual(r["blank_estimate"]["net_area_mm2"],
                               W * H - (4 - math.pi) * R * R - math.pi * 100 ** 2, delta=5)
        self.assertEqual(r["warnings"], [])

    def test_many_shuffled_segments(self):
        n = 360
        pts = [(100 * math.cos(2 * math.pi * i / n), 100 * math.sin(2 * math.pi * i / n)) for i in range(n)]
        order = list(range(n))
        random.Random(1).shuffle(order)
        doc = ezdxf.new("R2010")
        m = doc.modelspace()
        for k, i in enumerate(order):
            a, b = pts[i], pts[(i + 1) % n]
            m.add_line(*((b, a) if k % 2 else (a, b)))
        r, _ = dxfi.inspect(self.save(doc, "poly.dxf"), 5)
        self.assertEqual(r["closed_contours"][0]["type"], f"CHAIN({n})")
        # площадь правильного n-угольника n/2·R²·sin(2π/n)
        self.assertAlmostEqual(r["closed_contours"][0]["area_mm2"], n / 2 * 100 ** 2 * math.sin(2 * math.pi / n),
                               delta=0.05)

    def test_gap_is_reported_and_tol_closes_it(self):
        doc = ezdxf.new("R2010")
        m = doc.modelspace()
        m.add_line((0, 0), (1000, 0))
        m.add_line((1000, 0), (1000, 500))
        m.add_line((1000, 500), (0, 500))
        m.add_line((0, 500), (0, 1))  # разрыв 1 мм
        m.add_circle((200, 250), 50)
        path = self.save(doc, "gap.dxf")
        r, _ = dxfi.inspect(path, 5)
        self.assertEqual(r["open_chains"]["count"], 1)
        self.assertAlmostEqual(r["open_chains"]["length_mm"], 2999.0, places=1)
        self.assertTrue(any("Незамкнутые контуры" in w for w in r["warnings"]))
        self.assertFalse(r["blank_estimate"]["reliable"])
        self.assertIn("НЕДОСТОВЕРНА", r["blank_estimate"]["note"])
        r, _ = dxfi.inspect(path, 5, tol=1.5)
        self.assertTrue(r["blank_estimate"]["reliable"])
        self.assertAlmostEqual(r["blank_estimate"]["outer_area_mm2"], 500000, delta=300)  # замыкающая хорда 1 мм

    def test_layers_filter_and_cut_length_note(self):
        path = self.rect_with_holes(bend_line=True)
        r, _ = dxfi.inspect(path, 5)
        # 3000 + 2π·50 + 2π·30 + 500 (линия гиба) = 4002.65
        self.assertAlmostEqual(r["cut_length_total_mm"], 4002.7, delta=0.1)
        self.assertIn("ВСЕЙ геометрии", r["cut_length_note"])
        self.assertEqual(r["open_chains"]["count"], 1)
        r, _ = dxfi.inspect(path, 5, layers=["0"])
        self.assertAlmostEqual(r["cut_length_total_mm"], 3502.7, delta=0.1)
        self.assertEqual(r["open_chains"]["count"], 0)
        self.assertAlmostEqual(r["blank_estimate"]["net_area_mm2"], 489319, delta=1)

    def test_cli_tol_layers_json(self):
        path = self.rect_with_holes(bend_line=True)
        out = self.tmp / "r.json"
        buf = io.StringIO()
        with redirect_stdout(buf):
            rc = dxfi.main([path, "--layers", " 0 ", "--tol", "0.05", "--json", str(out)])
        self.assertEqual(rc, 0)
        self.assertIn("Заготовка (оценка)", buf.getvalue())
        data = json.loads(out.read_text(encoding="utf-8"))
        self.assertEqual(data["layers_filter"], ["0"])
        self.assertEqual(data["chain_tol_mm"], 0.05)
        self.assertAlmostEqual(data["cut_length_total_mm"], 3502.7, delta=0.1)


# ---------------------------------------------------------------------------------------------
class Unfold(unittest.TestCase):
    def run_unfold(self, *args):
        return subprocess.run([sys.executable, str(UNFOLD), *args], capture_output=True, text=True, encoding="utf-8")

    def test_outer_rejected_for_cone_gore_saddle(self):
        for args in (("cone", "--D", "1000", "--d", "500", "--H", "500", "--outer", "1016", "--t", "8", "--k", "0.5"),
                     ("gore", "--r", "159.5", "--R", "480", "--angle", "90", "--n", "3", "--t", "8"),
                     ("saddle", "--rb", "54", "--Rh", "159", "--k", "0.5")):
            p = self.run_unfold(*args)
            self.assertNotEqual(p.returncode, 0, args)
            self.assertIn("только в mitre", p.stderr, args)
            self.assertIn("нейтральн", p.stderr, args)

    def test_mitre_outer_converted(self):
        # D = 219 − 2·8 + 2·0.5·8 = 211
        p = self.run_unfold("mitre", "--outer", "219", "--t", "8", "--k", "0.5", "--L0", "300", "--beta", "22.5")
        self.assertEqual(p.returncode, 0, p.stderr)
        self.assertIn("Нейтральный диаметр D = 211.00 мм", p.stdout)

    def test_mitre_conflicting_or_partial_inputs(self):
        p = self.run_unfold("mitre", "--D", "211", "--outer", "219", "--t", "8", "--k", "0.5", "--L0", "300",
                            "--beta", "22.5")
        self.assertNotEqual(p.returncode, 0)
        p = self.run_unfold("mitre", "--outer", "219", "--t", "8", "--L0", "300", "--beta", "22.5")
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("--outer, --t, --k", p.stderr)


# ---------------------------------------------------------------------------------------------
def cost_data(labor_cat="labor", base="labor", kind="assumed"):
    return {"items": [
        {"name": "a", "category": " Material ", "kind": "CALC", "qty": 100, "unit_price": 10, "waste": 0.1,
         "source": "x", "uncertainty": 0.1},
        {"name": "b", "category": labor_cat, "kind": kind, "qty": 10, "unit_price": 100, "source": "x",
         "uncertainty": 0.2},
        {"name": "c", "category": "Purchased", "kind": "market", "price_kind": " Market", "qty": 2,
         "unit_price": 50, "source": "x", "uncertainty": 0.05}],
        "overheads": [{"name": "oh", "base": base, "percent": 100, "source": "x"},
                      {"name": "mat", "base": "Materials", "percent": 10, "source": "x"}],
        "contingency_percent": 0, "margin_percent": 10, "vat_percent": 20}


class CostNormalization(unittest.TestCase):
    def test_case_insensitive_categories_match_lowercase(self):
        # прямые = 1100 + 1000 + 100 = 2200; накладные 100 % от labor 1000 + 10 % от материалов (1100 + 100) = 120
        ref = cost.compute(cost_data())
        for cat, base in (("Labor", "LABOR"), (" labor ", "Labor ")):
            r = cost.compute(cost_data(cat, base))
            self.assertAlmostEqual(r["by_cat"]["labor"], 1000)
            self.assertAlmostEqual(r["overheads"][0]["base_value"], 1000)
            self.assertAlmostEqual(r["overheads"][1]["base_value"], 1200)
            self.assertAlmostEqual(r["cost_price"], ref["cost_price"])
            self.assertAlmostEqual(r["cost_price"], 2200 + 1000 + 120)
            self.assertEqual(r["warnings"], [])
        self.assertEqual({i["category"] for i in ref["items"]}, {"material", "labor", "purchased"})
        self.assertEqual([i["kind"] for i in ref["items"]], ["calc", "assumed", "market"])
        self.assertEqual(ref["items"][2]["price_kind"], "market")

    def test_unknown_category_and_kind_warn(self):
        r = cost.compute(cost_data("Labour", kind="guess"))
        b = r["items"][1]
        self.assertEqual((b["category"], b["kind"]), ("other", "assumed"))
        self.assertAlmostEqual(r["by_cat"]["other"], 1000)
        self.assertAlmostEqual(r["overheads"][0]["base_value"], 0)
        self.assertTrue(any("Labour" in w and "other" in w for w in r["warnings"]))
        self.assertTrue(any("guess" in w for w in r["warnings"]))
        md = cost.to_markdown({"title": "t"}, r)
        self.assertIn("| b |", md.split("## ДОПУЩЕНИЯ")[1])  # позиция видна в таблице, а не только в итогах

    def test_unknown_overhead_base_is_error(self):
        with self.assertRaises(SystemExit) as cm:
            cost.compute(cost_data(base="fot"))
        self.assertIn("неизвестная база", str(cm.exception))
        self.assertIn("labor | materials | direct", str(cm.exception))

    @unittest.skipIf(openpyxl is None, "нужен openpyxl")
    def test_xlsx_writes_normalized_categories(self):
        r = cost.compute(cost_data("Labor", "LABOR"))
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "e.xlsx")
            cost.to_xlsx({"title": "t"}, r, path)
            ws = openpyxl.load_workbook(path).active
            rows = list(ws.iter_rows(min_row=3, max_row=2 + len(r["items"]), values_only=True))
            cats = [row[3] for row in rows]
            self.assertEqual(cats, ["material", "labor", "purchased"])
            # SUMIF в Excel без учёта регистра: сумма по "labor" из столбца D = база накладных в Python
            excel_labor = sum(row[4] * (1 + (row[6] or 0)) * row[7] for row in rows if str(row[3]).lower() == "labor")
            self.assertAlmostEqual(excel_labor, r["overheads"][0]["base_value"])
            formulas = [c.value for c in ws["I"] if isinstance(c.value, str) and "SUMIF" in c.value]
            self.assertTrue(any('"labor"' in f for f in formulas))


# ---------------------------------------------------------------------------------------------
HOT_OK = {"G": 5.0, "cp": 4.19, "t_in": 80, "t_out": 59.27}   # Φ_г = 5·4.19·20.73 = 434.29 кВт
COLD_OK = {"G": 9.44, "cp": 2.3, "t_in": 5, "t_out": 25}      # Φ_х = 9.44·2.3·20 = 434.24 кВт


def check_ok(rep, prefix: str) -> bool:
    return next(c for c in rep.checks if c.text.startswith(prefix)).ok


class HeatBalance(unittest.TestCase):
    def test_consistent_balance_passes(self):
        r = hx.heat_balance({"hot": HOT_OK, "cold": COLD_OK})
        self.assertAlmostEqual(step_value(r, "Φ"), 434.3, places=1)
        self.assertEqual(r.verdict(), "ВЫПОЛНЕНО")

    def test_hot_side_sign_error_fails(self):
        hot = {**HOT_OK, "t_in": 59.27, "t_out": 80}  # вход и выход перепутаны → Φ_г < 0
        r = hx.heat_balance({"hot": hot, "cold": COLD_OK})
        self.assertFalse(check_ok(r, "Φ_г"))
        self.assertIn("проверьте вход/выход", next(c for c in r.checks if c.text.startswith("Φ_г")).note)
        self.assertFalse(check_ok(r, "Небаланс"))  # |Φ_г − Φ_х|/max = 2
        self.assertEqual(r.verdict(), "НЕ ВЫПОЛНЕНО")

    def test_both_sides_reversed_fail(self):
        hot = {**HOT_OK, "t_in": 59.27, "t_out": 80}
        cold = {**COLD_OK, "t_in": 25, "t_out": 5}
        r = hx.heat_balance({"hot": hot, "cold": cold})
        self.assertFalse(check_ok(r, "Φ_г"))
        self.assertFalse(check_ok(r, "Φ_х"))
        self.assertEqual(r.verdict(), "НЕ ВЫПОЛНЕНО")

    def test_two_unknowns_warn(self):
        r = hx.heat_balance({"hot": {"G": None, "cp": 4.19, "t_in": 80, "t_out": None}, "cold": COLD_OK})
        self.assertTrue(any("не заданы G, t_out" in w for w in r.warnings))

    def test_unknown_flow_with_reversed_temperatures_fails(self):
        r = hx.heat_balance({"hot": {"G": None, "cp": 4.19, "t_in": 60, "t_out": 80}, "cold": COLD_OK})
        self.assertEqual(r.verdict(), "НЕ ВЫПОЛНЕНО")
        with self.assertRaises(KeyError):
            step_value(r, "G")

    def test_single_unknowns(self):
        # G_г = 434.24/(4.19·20.73) = 4.9995; c_p = 434.24/(5·20.73) = 4.1895; t_вх = 59.27 + 434.24/(5·4.19) = 79.998
        r = hx.heat_balance({"hot": {**HOT_OK, "G": None}, "cold": COLD_OK})
        self.assertAlmostEqual(step_value(r, "G"), 5.0, places=2)
        r = hx.heat_balance({"hot": {**HOT_OK, "cp": None}, "cold": COLD_OK})
        self.assertAlmostEqual(step_value(r, "c_p"), 4.19, places=2)
        r = hx.heat_balance({"hot": {**HOT_OK, "t_in": None}, "cold": COLD_OK})
        self.assertAlmostEqual(step_value(r, "t_вх"), 80.0, places=1)
        r = hx.heat_balance({"hot": HOT_OK, "cold": {**COLD_OK, "t_out": None}})
        self.assertAlmostEqual(step_value(r, "t_вых"), 25.0, places=1)

    def test_no_full_side_exits(self):
        with self.assertRaises(SystemExit):
            hx.heat_balance({"hot": {"G": 5.0, "t_in": 80}, "cold": {"G": 9.44, "t_in": 5}})


# ---------------------------------------------------------------------------------------------
class BendB313(unittest.TestCase):
    EXACT = {"P": 10, "S": 138, "E": 1, "W": 1, "Y": 0.4, "R": 657, "D": 219, "c": 1.5}

    def test_eq_3c(self):
        # R1/D = 3: I_инт = 11/10 = 1.1; I_экстр = 13/14 = 0.92857
        # t_инт = 10·219/(2·(138/1.1 + 10·0.4)) = 8.4586; t_экстр = 10·219/(2·(138/0.92857 + 4)) = 7.1749
        r = pipe.bend_b31_3(dict(self.EXACT))
        self.assertAlmostEqual(step_value(r, "I_инт"), 1.1, places=4)
        self.assertAlmostEqual(step_value(r, "t_инт"), 8.459, places=3)
        self.assertAlmostEqual(step_value(r, "t_экстр"), 7.175, places=3)
        self.assertAlmostEqual(step_value(r, "t_m.инт"), 9.959, places=3)
        self.assertAlmostEqual(step_value(r, "t_m.экстр"), 8.675, places=3)
        self.assertFalse(any("ПРИБЛИЖЁННО" in w for w in r.warnings))

    def test_fallback_is_labelled(self):
        r = pipe.bend_b31_3({"t_straight": 10, "R": 657, "D": 219})
        self.assertAlmostEqual(step_value(r, "t_инт"), 11.0, places=3)
        self.assertTrue(any("приближённо (t·I), сверить с eq. 3c" in s.title for s in r.steps))
        self.assertTrue(any("ПРИБЛИЖЁННО" in w for w in r.warnings))

    def test_missing_inputs_exit(self):
        with self.assertRaises(SystemExit):
            pipe.bend_b31_3({"P": 10, "R": 657, "D": 219})
        with self.assertRaises(SystemExit):
            pipe.bend_b31_3({**self.EXACT, "R": 100})  # R1 < D/2


if __name__ == "__main__":
    unittest.main(verbosity=1)
