#!/usr/bin/env python3
"""Себестоимость и цена оборудования с разделением данных по происхождению.

Запуск:
    python cost_calc.py input.json [--out report.md] [--xlsx estimate.xlsx]

Каждая позиция обязана иметь:
  kind        calc     — РАСЧЁТНЫЕ ДАННЫЕ (масса из расчёта/модели, нормативная трудоёмкость по техпроцессу)
              market   — РЫНОЧНЫЕ ДАННЫЕ (КП поставщика, прайс, счёт; с датой и источником)
              assumed  — ДОПУЩЕНИЕ (экспертная оценка, аналог, укрупнённый норматив)
  category    material | purchased | labor | service | other
  source      откуда взято значение (документ, поставщик, дата)
  uncertainty относительная неопределённость цены/количества (0.1 = ±10 %)

{
  "title": "Фильтр-сепаратор ФС-1000-6,3",
  "currency": "руб.",
  "date": "2026-10-03",
  "items": [
    {"group": "Металлопрокат", "name": "Лист 09Г2С s=22", "category": "material", "kind": "calc",
     "qty": 2102, "unit": "кг", "waste": 0.12, "unit_price": 115, "price_kind": "market",
     "source": "масса — vessel_calc.py; цена — прайс металлобазы от 2026-09-20", "uncertainty": 0.08},
    ...
  ],
  "overheads": [{"name": "Цеховые накладные", "base": "labor", "percent": 120, "source": "учётная политика"}],
  "contingency_percent": 5, "margin_percent": 20,
  "vat_percent": {"value": 22, "source": "НК РФ ст. 164 (ред. с 01.01.2026) — проверить"}
}

price_kind (необязательно) — происхождение цены, если оно отличается от происхождения количества.
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict

KIND_TITLE = {"calc": "РАСЧЁТНЫЕ ДАННЫЕ", "market": "РЫНОЧНЫЕ ДАННЫЕ", "assumed": "ДОПУЩЕНИЯ"}


def val(x, default=None):
    if isinstance(x, dict):
        return x.get("value", default)
    return default if x is None else x


def src(x):
    return x.get("source", "ИСТОЧНИК НЕ УКАЗАН") if isinstance(x, dict) else "ИСТОЧНИК НЕ УКАЗАН"


def money(x: float) -> str:
    return f"{x:,.0f}".replace(",", " ")


def compute(data: dict) -> dict:
    items = []
    warnings = []
    for it in data.get("items", []):
        qty = float(it["qty"])
        waste = float(it.get("waste", 0.0))
        price = float(it["unit_price"])
        cost = qty * (1 + waste) * price
        unc = float(it.get("uncertainty", 0.0))
        if "uncertainty" not in it:
            warnings.append(f"«{it['name']}»: неопределённость не задана (принято 0).")
        if not it.get("source"):
            warnings.append(f"«{it['name']}»: не указан источник.")
        if it.get("kind") not in KIND_TITLE:
            warnings.append(f"«{it['name']}»: kind должен быть calc|market|assumed.")
        items.append({**it, "cost": cost, "low": cost * (1 - unc), "high": cost * (1 + unc), "unc": unc})

    by_cat = defaultdict(float)
    for it in items:
        by_cat[it.get("category", "other")] += it["cost"]
    direct = sum(it["cost"] for it in items)
    bases = {"labor": by_cat["labor"], "materials": by_cat["material"] + by_cat["purchased"], "direct": direct}

    overheads = []
    for oh in data.get("overheads", []):
        base = bases[oh["base"]]
        pct = float(val(oh.get("percent"), oh.get("percent")))
        overheads.append({**oh, "amount": base * pct / 100, "base_value": base, "pct": pct})
        if not oh.get("source"):
            warnings.append(f"Накладные «{oh['name']}»: не указан источник ставки.")
    cost_price = direct + sum(o["amount"] for o in overheads)
    cont_pct = float(val(data.get("contingency_percent"), 0.0))
    contingency = cost_price * cont_pct / 100
    full_cost = cost_price + contingency
    margin_pct = float(val(data.get("margin_percent"), 0.0))
    price_wo_vat = full_cost * (1 + margin_pct / 100)
    vat_pct = val(data.get("vat_percent"))
    if vat_pct is None:
        warnings.append("Ставка НДС не задана — цена с НДС не рассчитана.")
    vat = price_wo_vat * float(vat_pct) / 100 if vat_pct is not None else 0.0

    # неопределённость: независимые позиции -> корень из суммы квадратов; худший случай -> сумма
    import math
    sigma = math.sqrt(sum((it["cost"] * it["unc"]) ** 2 for it in items))
    worst = sum(it["cost"] * it["unc"] for it in items)
    contrib = sorted(items, key=lambda i: i["cost"] * i["unc"], reverse=True)[:5]
    return {
        "items": items, "by_cat": dict(by_cat), "direct": direct, "overheads": overheads,
        "cost_price": cost_price, "cont_pct": cont_pct, "contingency": contingency, "full_cost": full_cost,
        "margin_pct": margin_pct, "price_wo_vat": price_wo_vat, "vat_pct": vat_pct, "vat": vat,
        "price_with_vat": price_wo_vat + vat, "sigma": sigma, "worst": worst, "contrib": contrib,
        "warnings": warnings,
    }


def to_markdown(data: dict, r: dict) -> str:
    cur = data.get("currency", "руб.")
    out = [f"# Расчёт себестоимости и цены — {data.get('title', '')}", "",
           f"Дата расчёта: {data.get('date', 'не указана')}. Валюта: {cur}. Цены без НДС, если не указано иное.", ""]
    for kind, title in KIND_TITLE.items():
        rows = [i for i in r["items"] if i.get("kind") == kind]
        if not rows:
            continue
        out += [f"## {title}", "", "| Группа | Позиция | Кол-во | Ед. | Отход | Цена ед. | Происх. цены | Сумма | ± | Источник |",
                "|---|---|---|---|---|---|---|---|---|---|"]
        for i in rows:
            out.append(f"| {i.get('group', '')} | {i['name']} | {i['qty']:g} | {i.get('unit', '')} | "
                       f"{i.get('waste', 0) * 100:.0f} % | {money(i['unit_price'])} | {i.get('price_kind', kind)} | "
                       f"{money(i['cost'])} | {i['unc'] * 100:.0f} % | {i.get('source', '—')} |")
        out += [f"| | **Итого {title.lower()}** | | | | | | **{money(sum(i['cost'] for i in rows))}** | | |", ""]
    cat_names = {"material": "Материалы", "purchased": "Покупные изделия", "labor": "Трудоёмкость (ФОТ)",
                 "service": "Услуги (НК, испытания, покраска на стороне и т.п.)", "other": "Прочее"}
    out += ["## Сводка", "", "| Статья | Сумма |", "|---|---|"]
    for k, v in r["by_cat"].items():
        out.append(f"| {cat_names.get(k, k)} | {money(v)} |")
    out.append(f"| **Прямые затраты** | **{money(r['direct'])}** |")
    for o in r["overheads"]:
        out.append(f"| {o['name']} ({o['pct']:g} % от «{o['base']}» = {money(o['base_value'])}; {o.get('source', 'источник не указан')}) | {money(o['amount'])} |")
    out += [f"| **Производственная себестоимость** | **{money(r['cost_price'])}** |",
            f"| Резерв на непредвиденные расходы ({r['cont_pct']:g} %) | {money(r['contingency'])} |",
            f"| **Полная себестоимость** | **{money(r['full_cost'])}** |",
            f"| Прибыль ({r['margin_pct']:g} %) | {money(r['price_wo_vat'] - r['full_cost'])} |",
            f"| **Цена без НДС** | **{money(r['price_wo_vat'])}** |"]
    if r["vat_pct"] is not None:
        out += [f"| НДС ({r['vat_pct']:g} %; {src(data.get('vat_percent'))}) | {money(r['vat'])} |",
                f"| **Цена с НДС** | **{money(r['price_with_vat'])}** |"]
    out += ["", "## НЕОПРЕДЕЛЁННОСТЬ", "",
            f"- Прямые затраты: {money(r['direct'])} ± {money(r['sigma'])} (независимые позиции, СКО) "
            f"/ ± {money(r['worst'])} (худший случай, все отклонения в одну сторону).",
            f"- Диапазон цены без НДС при тех же наценках: {money(r['price_wo_vat'] * (1 - r['sigma'] / r['direct']))} — "
            f"{money(r['price_wo_vat'] * (1 + r['sigma'] / r['direct']))}." if r["direct"] else "",
            "- Наибольший вклад в неопределённость:"]
    for i in r["contrib"]:
        out.append(f"  - {i['name']}: ±{money(i['cost'] * i['unc'])} ({i.get('kind')}, {i.get('source', '—')})")
    out += ["", "Чтобы сузить диапазон — запросить КП по позициям из списка выше и уточнить трудоёмкость по техпроцессу.", ""]
    if r["warnings"]:
        out += ["## Предупреждения", ""] + [f"- ⚠️ {w}" for w in r["warnings"]] + [""]
    out += ["---", "_Расчёт для внутреннего использования. Рыночные цены действительны на дату источника; "
            "перед выдачей КП проверить актуальность цен, курс валют и ставку НДС._"]
    return "\n".join(x for x in out if x is not None)


def to_xlsx(data: dict, r: dict, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Смета"
    head = ["Группа", "Позиция", "Тип данных", "Категория", "Кол-во", "Ед.", "Отход, доля", "Цена ед.",
            "Сумма", "Неопр., доля", "Источник"]
    ws.append([f"Расчёт себестоимости — {data.get('title', '')}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(head)
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
    first = 3
    for n, it in enumerate(r["items"]):
        row = first + n
        ws.append([it.get("group", ""), it["name"], KIND_TITLE.get(it.get("kind"), it.get("kind")),
                   it.get("category", ""), it["qty"], it.get("unit", ""), it.get("waste", 0), it["unit_price"],
                   f"=E{row}*(1+G{row})*H{row}", it["unc"], it.get("source", "")])
    last = first + len(r["items"]) - 1
    row = last + 2
    ws.cell(row, 8, "Прямые затраты").font = Font(bold=True)
    ws.cell(row, 9, f"=SUM(I{first}:I{last})").font = Font(bold=True)
    direct_ref = f"I{row}"
    oh_refs = []
    for o in r["overheads"]:
        row += 1
        base = {"labor": f'SUMIF(D{first}:D{last},"labor",I{first}:I{last})',
                "materials": f'SUMIF(D{first}:D{last},"material",I{first}:I{last})+SUMIF(D{first}:D{last},"purchased",I{first}:I{last})',
                "direct": direct_ref}[o["base"]]
        ws.cell(row, 7, o["pct"] / 100)
        ws.cell(row, 8, f"{o['name']} (от {o['base']})")
        ws.cell(row, 9, f"=({base})*G{row}")
        oh_refs.append(f"I{row}")
    row += 1
    ws.cell(row, 8, "Производственная себестоимость").font = Font(bold=True)
    ws.cell(row, 9, "=" + "+".join([direct_ref] + oh_refs)).font = Font(bold=True)
    cp = f"I{row}"
    row += 1
    ws.cell(row, 7, r["cont_pct"] / 100)
    ws.cell(row, 8, "Резерв")
    ws.cell(row, 9, f"={cp}*G{row}")
    cont = f"I{row}"
    row += 1
    ws.cell(row, 8, "Полная себестоимость").font = Font(bold=True)
    ws.cell(row, 9, f"={cp}+{cont}").font = Font(bold=True)
    fc = f"I{row}"
    row += 1
    ws.cell(row, 7, r["margin_pct"] / 100)
    ws.cell(row, 8, "Цена без НДС")
    ws.cell(row, 9, f"={fc}*(1+G{row})").font = Font(bold=True)
    pw = f"I{row}"
    if r["vat_pct"] is not None:
        row += 1
        ws.cell(row, 7, float(r["vat_pct"]) / 100)
        ws.cell(row, 8, "Цена с НДС")
        ws.cell(row, 9, f"={pw}*(1+G{row})").font = Font(bold=True)
    for col, w in zip("ABCDEFGHIJK", (16, 40, 18, 12, 10, 6, 10, 14, 16, 10, 50)):
        ws.column_dimensions[col].width = w
    for rr in ws.iter_rows(min_row=first, max_row=row, min_col=8, max_col=9):
        for c in rr:
            c.number_format = '#,##0'
    wb.save(path)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input")
    ap.add_argument("--out")
    ap.add_argument("--xlsx")
    a = ap.parse_args(argv)
    with open(a.input, encoding="utf-8") as f:
        data = json.load(f)
    r = compute(data)
    md = to_markdown(data, r)
    if a.out:
        with open(a.out, "w", encoding="utf-8") as f:
            f.write(md)
        print(f"Отчёт сохранён: {a.out}")
    else:
        sys.stdout.reconfigure(encoding="utf-8")
        print(md)
    if a.xlsx:
        to_xlsx(data, r, a.xlsx)
        print(f"Excel сохранён: {a.xlsx} (формулы пересчитываются в Excel/LibreOffice)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
